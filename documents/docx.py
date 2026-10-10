from io import BytesIO
from copy import deepcopy
import re
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document
from docxtpl import DocxTemplate
from jinja2 import Environment, nodes
from jinja2.visitor import NodeVisitor


class DocxTemplateError(ValueError):
    """The uploaded DOCX or its variable contract is invalid."""


_VARIABLE_NAME = re.compile(r'^[A-Za-z_][A-Za-z0-9_]*$')
_SUPPORTED_TYPES = {'string', 'integer', 'number', 'boolean', 'object', 'list'}
_WORD_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _normalise_field(name, field):
    if not isinstance(field, dict):
        raise DocxTemplateError(f'Le contrat de {name} doit être un objet.')

    data_type = field.get('data_type', 'string')
    if data_type not in _SUPPORTED_TYPES:
        raise DocxTemplateError(f'Type non pris en charge pour {name} : {data_type}.')

    required = field.get('is_required', False)
    if not isinstance(required, bool):
        raise DocxTemplateError(f'Le caractère obligatoire de {name} doit être booléen.')

    normalized = {**field, 'data_type': data_type, 'is_required': required}
    properties = field.get('properties', {})
    if not isinstance(properties, dict):
        raise DocxTemplateError(f'Les propriétés de {name} doivent être un objet.')
    normalized['properties'] = {}
    for property_name, property_specification in properties.items():
        if not isinstance(property_name, str) or not _VARIABLE_NAME.fullmatch(property_name):
            raise DocxTemplateError(f'Nom de propriété invalide pour {name}.')
        normalized['properties'][property_name] = _normalise_field(
            f'{name}.{property_name}',
            property_specification,
        )

    if data_type == 'list':
        allow_empty = field.get('allow_empty')
        if not isinstance(allow_empty, bool):
            raise DocxTemplateError(f'La règle allow_empty doit être définie pour {name}.')
        items = field.get('items')
        if not isinstance(items, dict):
            raise DocxTemplateError(f'Le schéma des éléments de {name} est obligatoire.')
        normalized['allow_empty'] = allow_empty
        normalized['items'] = _normalise_field(f'{name}[]', items)

    return normalized


def _normalise_contract(variables):
    if not isinstance(variables, list):
        raise DocxTemplateError('Le contrat de variables doit être une liste.')

    contract = {}
    for index, variable in enumerate(variables):
        if not isinstance(variable, dict):
            raise DocxTemplateError(f'La variable {index} doit être un objet.')

        name = variable.get('name')
        if not isinstance(name, str) or not _VARIABLE_NAME.fullmatch(name):
            raise DocxTemplateError(f'Nom de variable invalide à la position {index}.')
        if name in contract:
            raise DocxTemplateError(f'Variable déclarée plusieurs fois : {name}.')

        contract[name] = _normalise_field(name, variable)

    return contract


def _load_template(template_bytes):
    if not template_bytes:
        raise DocxTemplateError('Le fichier template DOCX est vide.')

    try:
        template = DocxTemplate(BytesIO(template_bytes))
        template.init_docx()
        return template
    except DocxTemplateError:
        raise
    except Exception as exc:
        raise DocxTemplateError(
            'Le fichier template DOCX est illisible ou invalide.'
        ) from exc


def _jinja_source(template_bytes):
    paragraphs = []
    try:
        with ZipFile(BytesIO(template_bytes)) as package:
            document_parts = [
                name for name in package.namelist()
                if name == 'word/document.xml'
                or re.fullmatch(r'word/(?:header|footer)\d+\.xml', name)
            ]
            for part_name in document_parts:
                root = ElementTree.fromstring(package.read(part_name))
                for paragraph in root.iter(f'{_WORD_NS}p'):
                    paragraphs.append(''.join(
                        text.text or '' for text in paragraph.iter(f'{_WORD_NS}t')
                    ))
    except Exception as exc:
        raise DocxTemplateError('Le contenu XML du template DOCX est invalide.') from exc

    source = '\n'.join(paragraphs)
    return re.sub(r'{%(-?)(?:p|tr|tc|r)\b', r'{%\1', source)


def _expression_path(expression):
    if isinstance(expression, nodes.Name):
        return expression.name
    if isinstance(expression, nodes.Getattr):
        base = _expression_path(expression.node)
        return f'{base}.{expression.attr}' if base else None
    if isinstance(expression, nodes.Getitem):
        base = _expression_path(expression.node)
        if isinstance(expression.arg, nodes.Const) and isinstance(expression.arg.value, str):
            return f'{base}.{expression.arg.value}' if base else None
    return None


class _AttributePathVisitor(NodeVisitor):
    def __init__(self):
        self.paths = set()
        self.loop_paths = {}

    def visit_For(self, node, *args, **kwargs):
        source_path = _expression_path(node.iter)
        target_names = [node.target, *node.target.find_all(nodes.Name)]
        names = [target.name for target in target_names if isinstance(target, nodes.Name)]
        previous = {name: self.loop_paths.get(name) for name in names}
        if source_path:
            for name in names:
                self.loop_paths[name] = f'{source_path}[]'
        for child in [*node.body, *node.else_]:
            self.visit(child)
        for name, prior_path in previous.items():
            if prior_path is None:
                self.loop_paths.pop(name, None)
            else:
                self.loop_paths[name] = prior_path
        self.visit(node.iter)

    def visit_Getattr(self, node, *args, **kwargs):
        path = _expression_path(node)
        if path:
            root, separator, remainder = path.partition('.')
            loop_root = self.loop_paths.get(root)
            if loop_root:
                path = f'{loop_root}.{remainder}' if separator else loop_root
            self.paths.add(path)
        self.generic_visit(node)

    def visit_Getitem(self, node, *args, **kwargs):
        path = _expression_path(node)
        if path:
            root, separator, remainder = path.partition('.')
            loop_root = self.loop_paths.get(root)
            if loop_root:
                path = f'{loop_root}.{remainder}' if separator else loop_root
            self.paths.add(path)
        self.generic_visit(node)


def _path_is_declared(path, contract):
    parts = re.findall(r'[A-Za-z_][A-Za-z0-9_]*|\[\]', path)
    if not parts:
        return False
    specification = contract.get(parts.pop(0))
    if specification is None:
        return False

    for part in parts:
        if part == '[]':
            if specification['data_type'] != 'list':
                return False
            specification = specification['items']
        else:
            specification = specification['properties'].get(part)
            if specification is None:
                return False
    return True


def validate_docx_template(template_bytes, variables):
    contract = _normalise_contract(variables)
    template = _load_template(template_bytes)
    try:
        placeholders = template.get_undeclared_template_variables()
        syntax_tree = Environment().parse(_jinja_source(template_bytes))
    except Exception as exc:
        raise DocxTemplateError(
            'Les placeholders ou blocs du template sont mal formés.'
        ) from exc

    unknown = sorted(placeholders - contract.keys())
    visitor = _AttributePathVisitor()
    visitor.visit(syntax_tree)
    unknown.extend(
        path for path in visitor.paths
        if not _path_is_declared(path, contract)
    )
    if unknown:
        raise DocxTemplateError(
            f'Variables ou propriétés présentes dans le DOCX mais absentes du contrat : '
            f'{", ".join(sorted(set(unknown)))}.'
        )
    return contract


def _validate_value(name, specification, value):
    data_type = specification['data_type']
    type_checks = {
        'string': lambda item: isinstance(item, str),
        'integer': lambda item: isinstance(item, int) and not isinstance(item, bool),
        'number': lambda item: isinstance(item, (int, float)) and not isinstance(item, bool),
        'boolean': lambda item: isinstance(item, bool),
        'object': lambda item: isinstance(item, dict),
        'list': lambda item: isinstance(item, list),
    }
    if not type_checks[data_type](value):
        raise DocxTemplateError(
            f'Valeur incompatible avec le type {data_type} pour {name}.'
        )
    if data_type == 'list' and not value and not specification['allow_empty']:
        raise DocxTemplateError(f'La liste {name} ne peut pas être vide.')
    if data_type == 'object':
        for property_name, property_specification in specification['properties'].items():
            if property_name not in value or value[property_name] is None:
                if property_specification['is_required']:
                    raise DocxTemplateError(
                        f'Propriété obligatoire absente : {name}.{property_name}.'
                    )
                continue
            _validate_value(
                f'{name}.{property_name}',
                property_specification,
                value[property_name],
            )
    if data_type == 'list':
        for index, item in enumerate(value):
            _validate_value(f'{name}[{index}]', specification['items'], item)


def render_docx(template_bytes, variables, context):
    contract = validate_docx_template(template_bytes, variables)
    if not isinstance(context, dict):
        raise DocxTemplateError('Le contexte de génération doit être un objet.')

    render_context = deepcopy(context)
    for name, specification in contract.items():
        value = render_context.get(name)
        if name not in render_context or value is None or value == '':
            if specification['is_required']:
                raise DocxTemplateError(f'Variable obligatoire absente : {name}.')
            if specification.get('default_value') is not None:
                value = specification['default_value']
            elif specification['data_type'] == 'list':
                value = []
            elif specification['data_type'] == 'object':
                value = {}
            else:
                value = ''
            render_context[name] = value

        _validate_value(name, specification, value)

    template = _load_template(template_bytes)
    try:
        template.render(render_context)
        output = BytesIO()
        template.save(output)
        content = output.getvalue()
        Document(BytesIO(content))
        return content
    except DocxTemplateError:
        raise
    except Exception as exc:
        raise DocxTemplateError('Le rendu du template DOCX a échoué.') from exc
