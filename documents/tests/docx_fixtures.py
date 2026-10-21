from io import BytesIO
from pathlib import Path

from docx import Document


# Copie de travail du véritable modèle Word enregistré dans l'application
# (DocumentTemplate « Ordonnance standard », stocké dans le volume média).
# Le fichier source n'est jamais modifié : seules des copies sont utilisées
# pour les tests et les rendus.
REAL_TEMPLATE_PATH = (
    Path(__file__).resolve().parent / "fixtures" / "ordonnance_standard.docx"
)


def load_real_template_bytes():
    """Octets exacts du véritable modèle « Ordonnance standard »."""
    return REAL_TEMPLATE_PATH.read_bytes()


def make_prescription_docx(*, patient_name_variable="{{ patient.first_name }}"):
    document = Document()
    document.sections[0].header.paragraphs[0].text = (
        "Praticien : {{ dentist.first_name }} {{ dentist.last_name }}"
    )
    document.add_paragraph("{{ prescription_start }}")

    patient_line = document.add_paragraph("Patient : ")
    patient_line.add_run(patient_name_variable.split("{{")[0])
    variable = patient_name_variable.strip()
    if variable == "{{ patient.first_name }}":
        patient_line.add_run("{{ patient.first_")
        patient_line.add_run("name }}")
    else:
        patient_line.add_run(variable)
    patient_line.add_run(" {{ patient.last_name }}")

    table = document.add_table(rows=1, cols=4)
    table.rows[0].cells[0].text = "{{ medication.name }}"
    table.rows[0].cells[1].text = "{{ medication.dosage }}"
    table.rows[0].cells[2].text = "{{ medication.frequency }}"
    table.rows[0].cells[3].text = "{{ medication.notes }}"
    document.add_paragraph("{{ prescription_end }}")
    document.add_paragraph("Signature")
    document.sections[0].footer.paragraphs[0].text = (
        "{{ prescription.date }} — {{ prescription.notes }}"
    )

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_simple_docx(text="Modèle administratif"):
    document = Document()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_docx_with_text(text):
    document = Document()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
