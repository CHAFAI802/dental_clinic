import re
from copy import deepcopy
from dataclasses import dataclass
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
from zipfile import BadZipFile

from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from django.core.exceptions import ValidationError
from lxml.etree import XMLSyntaxError


_VARIABLE_PATTERN = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
_MEDICATION_FIELDS = {
    "medication.name",
    "medication.dosage",
    "medication.frequency",
    "medication.route",
    "medication.start_date",
    "medication.end_date",
    "medication.notes",
}
_SCALAR_FIELDS = {
    "patient_name",
    "document_date",
    "patient.first_name",
    "patient.middle_name",
    "patient.last_name",
    "patient.patient_code",
    "patient.birthdate",
    "patient.phone",
    "patient.address",
    "patient.city",
    "dentist.first_name",
    "dentist.last_name",
    "dentist.email",
    "dentist.phone",
    "prescription.date",
    "prescription.notes",
}
_ZONE_MARKERS = {"prescription_start", "prescription_end"}
# Bloc de prescription en texte libre : le vrai modèle « Ordonnance standard »
# (media/document_templates/ordonnance_*.docx) injecte la liste complète des
# médicaments dans un paragraphe, sans tableau répétable.
_BLOCK_FIELDS = {"prescription_details"}
_SUPPORTED_FIELDS = (
    _MEDICATION_FIELDS | _SCALAR_FIELDS | _ZONE_MARKERS | _BLOCK_FIELDS
)


@dataclass(frozen=True)
class DocxTemplateInfo:
    variables: tuple[str, ...]
    has_medication_rows: bool
    has_prescription_details: bool = False

    @property
    def expects_medications(self):
        """Le modèle exige au moins un médicament pour être rempli."""
        return self.has_medication_rows or self.has_prescription_details

    @property
    def supports_medications(self):
        """Le modèle sait restituer une liste de médicaments."""
        return self.has_medication_rows or self.has_prescription_details


def file_sha256(file_obj):
    file_obj.open("rb")
    digest = sha256()
    for chunk in file_obj.chunks():
        digest.update(chunk)
    return digest.hexdigest()


def inspect_docx_template(content):
    document = _open_document(content)
    found_variables = []
    story_text = []
    story_variables = []

    for paragraphs, tables in _document_stories(document):
        current_story_variables = []
        for paragraph in paragraphs:
            text = paragraph.text
            variables = _variables_in_text(text)
            if set(variables) & _MEDICATION_FIELDS and not _paragraph_in_table_row(
                paragraph
            ):
                raise ValidationError(
                    "Les variables medication.* doivent se trouver dans une ligne "
                    "de tableau répétable."
                )
            found_variables.extend(variables)
            current_story_variables.extend(variables)
            story_text.append(text)
        story_variables.append(current_story_variables)

    _validate_tokens(story_text, found_variables, story_variables)
    return DocxTemplateInfo(
        variables=tuple(dict.fromkeys(found_variables)),
        has_medication_rows=any(
            set(variable_list) & _MEDICATION_FIELDS
            for variable_list in story_variables
        ),
        has_prescription_details="prescription_details" in found_variables,
    )


def render_prescription_docx(
    content,
    *,
    patient,
    dentist,
    generated_at,
    notes,
    medications,
):
    info = inspect_docx_template(content)
    if info.expects_medications and not medications:
        raise ValidationError(
            "Ce modèle contient une zone de prescription, mais aucune donnée "
            "médicamenteuse n’a été sélectionnée."
        )

    document = _open_document(content)
    scalar_values = _scalar_values(
        patient, dentist, generated_at, notes, medications
    )
    for _, tables in _document_stories(document):
        for table in tables:
            _expand_medication_rows(table, medications, scalar_values)

    for paragraphs, _ in _document_stories(document):
        for paragraph in paragraphs:
            _replace_paragraph_tokens(paragraph, scalar_values)

    output = BytesIO()
    document.save(output)
    generated_content = output.getvalue()
    remaining = inspect_docx_template(generated_content)
    if remaining.variables:
        raise ValidationError(
            "Des variables ou des marqueurs restent non résolus dans le document généré."
        )
    return generated_content


def _open_document(content):
    if not content:
        raise ValidationError("Le fichier DOCX est vide.")
    try:
        return WordDocument(BytesIO(content))
    except (BadZipFile, PackageNotFoundError, XMLSyntaxError, ValueError, KeyError) as error:
        raise ValidationError(
            "Le fichier n’est pas un document DOCX valide ou lisible."
        ) from error


def _document_stories(document):
    yield _iter_document_paragraphs(document), _iter_document_tables(document)
    for section in document.sections:
        for part in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            yield _iter_document_paragraphs(part), _iter_document_tables(part)


def _iter_document_paragraphs(container):
    root = container._element.body if hasattr(container._element, "body") else container._element
    for child in root.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, container)
        elif isinstance(child, CT_Tbl):
            yield from _iter_table_paragraphs(Table(child, container))


def _iter_document_tables(container):
    for table in container.tables:
        yield from _iter_nested_tables(table)


def _iter_nested_tables(table):
    yield table
    for row in table.rows:
        for cell in row.cells:
            for nested_table in cell.tables:
                yield from _iter_nested_tables(nested_table)


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from _iter_cell_paragraphs(cell)


def _iter_cell_paragraphs(cell):
    yield from cell.paragraphs
    for nested_table in cell.tables:
        yield from _iter_table_paragraphs(nested_table)


def _variables_in_text(text):
    return [match.group(1).strip() for match in _VARIABLE_PATTERN.finditer(text)]


def _validate_tokens(story_text, found_variables, story_variables):
    unknown = sorted(set(found_variables) - _SUPPORTED_FIELDS)
    if unknown:
        raise ValidationError(
            f"Variables DOCX inconnues : {', '.join(unknown)}."
        )

    for text in story_text:
        stripped = _VARIABLE_PATTERN.sub("", text)
        if "{{" in stripped or "}}" in stripped:
            raise ValidationError(
                "Le modèle contient un marqueur de variable mal formé."
            )

    starts = found_variables.count("prescription_start")
    ends = found_variables.count("prescription_end")
    if starts != ends or starts > 1:
        raise ValidationError(
            "Les marqueurs prescription_start et prescription_end doivent être "
            "présents ensemble et chacun une seule fois."
        )
    if starts:
        marker_story = next(
            variables
            for variables in story_variables
            if "prescription_start" in variables or "prescription_end" in variables
        )
        if marker_story.index("prescription_start") > marker_story.index(
            "prescription_end"
        ):
            raise ValidationError(
                "Le marqueur prescription_start doit précéder prescription_end."
            )
        if any(field in found_variables for field in _MEDICATION_FIELDS):
            start = marker_story.index("prescription_start")
            end = marker_story.index("prescription_end")
            if not any(
                variable in _MEDICATION_FIELDS
                for variable in marker_story[start + 1 : end]
            ):
                raise ValidationError(
                    "La zone de prescription doit contenir les variables "
                    "medication.*."
                )
    elif any(field in found_variables for field in _MEDICATION_FIELDS):
        raise ValidationError(
            "Les modèles contenant des médicaments doivent délimiter leur zone "
            "avec prescription_start et prescription_end."
        )


def _paragraph_in_table_row(paragraph):
    row = _containing_row(paragraph)
    return row is not None


def _containing_row(paragraph):
    parent = paragraph._p.getparent()
    while parent is not None:
        if parent.tag.endswith("}tr"):
            return parent
        parent = parent.getparent()
    return None


def _expand_medication_rows(table, medications, scalar_values):
    rows = list(table.rows)
    for row in rows:
        row_paragraphs = [
            paragraph
            for cell in row.cells
            for paragraph in _iter_cell_paragraphs(cell)
        ]
        row_text = "\n".join(paragraph.text for paragraph in row_paragraphs)
        row_variables = set(_variables_in_text(row_text))
        if not row_variables & _MEDICATION_FIELDS:
            continue

        template_xml = deepcopy(row._tr)
        insertion_point = row._tr
        for medication in medications:
            clone_xml = deepcopy(template_xml)
            insertion_point.addnext(clone_xml)
            insertion_point = clone_xml
            clone_row = next(
                row_wrapper
                for row_wrapper in table.rows
                if row_wrapper._tr == clone_xml
            )
            _replace_row_tokens(
                clone_row,
                {
                    **scalar_values,
                    **_medication_values(medication),
                },
            )
        row._tr.getparent().remove(row._tr)


def _replace_row_tokens(row, values):
    for cell in row.cells:
        for paragraph in _iter_cell_paragraphs(cell):
            _replace_paragraph_tokens(paragraph, values)


def _replace_paragraph_tokens(paragraph, values):
    runs = paragraph.runs
    text = "".join(run.text for run in runs)
    matches = list(_VARIABLE_PATTERN.finditer(text))
    if not matches:
        return

    offsets = []
    offset = 0
    for run in runs:
        offsets.append(offset)
        offset += len(run.text)

    for match in reversed(matches):
        name = match.group(1).strip()
        if name in _ZONE_MARKERS:
            replacement = ""
        elif name in values:
            replacement = values[name]
        else:
            raise ValidationError(
                f"La variable « {name} » ne possède aucune valeur pour ce document."
            )
        start, end = match.span()
        intersecting_runs = [
            index
            for index, run_start in enumerate(offsets)
            if run_start < end and run_start + len(runs[index].text) > start
        ]
        first_run = intersecting_runs[0]
        last_run = intersecting_runs[-1]
        if first_run == last_run:
            run_start = offsets[first_run]
            local_start = start - run_start
            local_end = end - run_start
            runs[first_run].text = (
                runs[first_run].text[:local_start]
                + replacement
                + runs[first_run].text[local_end:]
            )
            continue

        first_local = start - offsets[first_run]
        runs[first_run].text = runs[first_run].text[:first_local] + replacement
        for index in range(first_run + 1, last_run):
            runs[index].text = ""
        last_local_end = end - offsets[last_run]
        runs[last_run].text = runs[last_run].text[last_local_end:]


def _scalar_values(patient, dentist, generated_at, notes, medications=()):
    return {
        "patient_name": _patient_name(patient),
        "document_date": _format_date(generated_at),
        "prescription_details": _prescription_details(medications),
        "patient.first_name": patient.first_name,
        "patient.middle_name": patient.middle_name,
        "patient.last_name": patient.last_name,
        "patient.patient_code": patient.patient_code,
        "patient.birthdate": _format_date(patient.birthdate),
        "patient.phone": patient.phone,
        "patient.address": patient.address,
        "patient.city": patient.city,
        "dentist.first_name": dentist.first_name,
        "dentist.last_name": dentist.last_name,
        "dentist.email": dentist.email,
        "dentist.phone": dentist.phone,
        "prescription.date": _format_date(generated_at),
        "prescription.notes": notes,
    }


def _patient_name(patient):
    return " ".join(
        part
        for part in (patient.first_name, patient.middle_name, patient.last_name)
        if part
    )


def _prescription_details(medications):
    """Rend la liste des médicaments dans un paragraphe texte libre.

    Chaque médicament produit un bloc séparé d'une ligne vide ; les sauts de
    ligne sont convertis en <w:br/> par python-docx, la mise en page du
    modèle (titre, signature, pied de page) n'est pas modifiée.
    """
    blocks = []
    for medication in medications:
        header = medication.name
        if medication.dosage:
            header = f"{header} — {medication.dosage}"
        lines = [header]
        if medication.frequency:
            lines.append(f"- Fréquence : {medication.frequency}")
        if medication.route:
            lines.append(f"- Voie : {medication.route}")
        if medication.start_date:
            lines.append(f"- Début : {_format_date(medication.start_date)}")
        if medication.end_date:
            lines.append(f"- Fin : {_format_date(medication.end_date)}")
        if medication.start_date and medication.end_date:
            duration = (medication.end_date - medication.start_date).days + 1
            lines.append(f"- Durée : {duration} jours")
        if medication.notes:
            lines.append(f"- Notes : {medication.notes}")
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def _medication_values(medication):
    return {
        "medication.name": medication.name,
        "medication.dosage": medication.dosage,
        "medication.frequency": medication.frequency,
        "medication.route": medication.route,
        "medication.start_date": _format_date(medication.start_date),
        "medication.end_date": _format_date(medication.end_date),
        "medication.notes": medication.notes,
    }


def _format_date(value):
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return ""
