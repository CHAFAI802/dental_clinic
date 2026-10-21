"""Test de bout en bout de la génération d'ordonnances.

Parcours réel couvert, du client API jusqu'au fichier DOCX produit :

    authentification du médecin
        -> sélection d'un patient autorisé
        -> récupération des données du patient (source métier : tables Patient)
        -> envoi des médicaments (contrat réel de l'API)
        -> appel du véritable endpoint generate-from-document-template
        -> validation du modèle actif ET approuvé
        -> génération du DOCX à partir du véritable modèle Word
        -> contrôle du contenu, de la mise en page et du modèle source

Le modèle utilisé est une copie de travail, octet pour octet, du véritable
modèle « Ordonnance standard » enregistré dans l'application (aucun modèle
artificiel n'est substitué au document réel).
"""

from datetime import date, timedelta
from hashlib import sha256
from io import BytesIO
from tempfile import TemporaryDirectory
from zipfile import is_zipfile

from docx import Document as WordDocument
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.utils import timezone
from rest_framework.test import APITestCase

from accounts.models import User
from appointments.models import Appointment
from documents.models import DocumentTemplate
from documents.tests.docx_fixtures import load_real_template_bytes
from patients.models import Medication, Patient
from prescriptions.models import Prescription


def _document_text(document):
    """Texte complet du DOCX : corps, tableaux, en-têtes et pieds de page."""
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    for section in document.sections:
        for part in (section.header, section.footer):
            parts.extend(paragraph.text for paragraph in part.paragraphs)
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
    return "\n".join(parts)


class RealTemplateEndToEndTests(APITestCase):
    def setUp(self):
        self.media_dir = TemporaryDirectory()
        self.addCleanup(self.media_dir.cleanup)
        self.settings_context = override_settings(MEDIA_ROOT=self.media_dir.name)
        self.settings_context.enable()
        self.addCleanup(self.settings_context.disable)

        self.dentist = User.objects.create_user(
            email="e2e-dentist@example.test",
            password="test-password",
            first_name="Nadia",
            last_name="Ben Salah",
            role=User.Role.DENTIST,
        )
        self.other_dentist = User.objects.create_user(
            email="e2e-other@example.test",
            password="test-password",
            first_name="Autre",
            last_name="Praticien",
            role=User.Role.DENTIST,
        )
        self.administrator = User.objects.create_user(
            email="e2e-admin@example.test",
            password="test-password",
            first_name="Admin",
            last_name="User",
            role=User.Role.ADMINISTRATOR,
        )
        self.patient = Patient.objects.create(
            first_name="Émilie",
            middle_name="Marie",
            last_name="Durand",
            birthdate=date(1990, 3, 15),
            phone="06 12 34 56 78",
            address="12 rue des Lilas",
            city="Lyon",
        )
        starts = timezone.now() + timedelta(days=1)
        Appointment.objects.create(
            patient=self.patient,
            practitioner=self.dentist,
            start_at=starts,
            end_at=starts + timedelta(minutes=30),
        )
        self.medications = [
            Medication.objects.create(
                patient=self.patient,
                prescribed_by=self.dentist,
                name="Amoxicilline",
                dosage="500 mg",
                frequency="3 fois par jour",
                route="orale",
                start_date=date(2026, 10, 15),
                end_date=date(2026, 10, 21),
                notes="Après le repas, cure de 7 jours",
            ),
            Medication.objects.create(
                patient=self.patient,
                prescribed_by=self.dentist,
                name="Paracétamol",
                dosage="1 g",
                frequency="si douleur",
                route="orale",
                start_date=date(2026, 10, 15),
                notes="Ne pas dépasser 3 g par jour — éviter l'ibuprofène",
            ),
        ]

        # Véritable modèle Word : copie de travail des octets d'origine.
        self.real_template_bytes = load_real_template_bytes()
        self.template = DocumentTemplate.objects.create(
            name="Ordonnance standard",
            file=SimpleUploadedFile(
                "ordonnance_standard.docx",
                self.real_template_bytes,
            ),
        )
        self.client.force_authenticate(self.administrator)
        approval = self.client.post(
            f"/api/document-templates/{self.template.pk}/approve/",
            format="json",
        )
        self.assertEqual(approval.status_code, 200)

    # ------------------------------------------------------------------
    # Parcours nominal
    # ------------------------------------------------------------------
    def test_end_to_end_generation_with_real_word_template(self):
        self.template.refresh_from_db()
        source_hash_before = sha256(self.template.file.read()).hexdigest()
        self.template.file.seek(0)

        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [medication.pk for medication in self.medications],
                "notes": "Contrôle post-opératoire",
            },
            format="json",
        )

        # 1. Réponse HTTP et enregistrement de l'artefact.
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["status"], Prescription.Status.GENERATED)
        self.assertTrue(response.data["generated_file"])
        self.assertEqual(len(response.data["filled_data"]["medications"]), 2)

        prescription = Prescription.objects.get(pk=response.data["id"])
        self.assertEqual(prescription.patient_id, self.patient.pk)
        self.assertEqual(prescription.dentist_id, self.dentist.pk)
        self.assertEqual(prescription.document_template_id, self.template.pk)
        self.assertIsNone(prescription.template_id)
        self.assertTrue(prescription.generated_file.name.startswith("prescriptions/generated/"))
        self.assertTrue(prescription.generated_file.name.endswith(".docx"))

        with prescription.generated_file.open("rb") as generated:
            generated_content = generated.read()
        self.assertTrue(is_zipfile(BytesIO(generated_content)))
        document = WordDocument(BytesIO(generated_content))
        body_text = _document_text(document)

        # 2. Injection des données patient depuis la source métier.
        self.assertIn("ORDONNANCE MÉDICALE", body_text)
        self.assertIn("Clinique dentaire", body_text)
        self.assertIn("Patient : Émilie Marie Durand", body_text)
        self.assertIn(
            f"Date : {prescription.generated_at.strftime('%d/%m/%Y')}",
            body_text,
        )

        # 3. Injection des prescriptions (bloc {{ prescription_details }}).
        self.assertIn("Amoxicilline — 500 mg", body_text)
        self.assertIn("- Fréquence : 3 fois par jour", body_text)
        self.assertIn("- Voie : orale", body_text)
        self.assertIn("- Début : 15/10/2026", body_text)
        self.assertIn("- Fin : 21/10/2026", body_text)
        self.assertIn("- Durée : 7 jours", body_text)
        self.assertIn("Paracétamol — 1 g", body_text)
        self.assertIn("si douleur", body_text)
        self.assertIn(
            "Ne pas dépasser 3 g par jour — éviter l'ibuprofène",
            body_text,
        )
        self.assertLess(
            body_text.index("Amoxicilline"),
            body_text.index("Paracétamol"),
            "L'ordre des prescriptions doit être respecté.",
        )

        # 4. Aucun marqueur technique ne subsiste (corps, tableaux, en-tête, pied).
        self.assertNotIn("{{", body_text)
        self.assertNotIn("}}", body_text)

        # 5. Conservation de la mise en page du modèle original.
        self.assertIn("Signature et cachet du praticien", body_text)
        self.assertEqual(
            [paragraph.text for paragraph in document.paragraphs][:4],
            [
                "ORDONNANCE MÉDICALE",
                "Clinique dentaire",
                f"Patient : Émilie Marie Durand",
                f"Date : {prescription.generated_at.strftime('%d/%m/%Y')}",
            ],
        )
        self.assertEqual(len(document.sections), 1)
        self.assertEqual(
            document.sections[0].header.paragraphs[0].text,
            WordDocument(BytesIO(self.real_template_bytes)).sections[0].header.paragraphs[0].text,
        )
        self.assertEqual(
            document.sections[0].footer.paragraphs[0].text,
            WordDocument(BytesIO(self.real_template_bytes)).sections[0].footer.paragraphs[0].text,
        )
        self.assertEqual(
            document.sections[0].bottom_margin,
            WordDocument(BytesIO(self.real_template_bytes)).sections[0].bottom_margin,
        )

        # 6. Le modèle source reste intact et l'artefact en est distinct.
        self.template.refresh_from_db()
        source_hash_after = sha256(self.template.file.read()).hexdigest()
        self.assertEqual(source_hash_before, source_hash_after)
        self.assertEqual(source_hash_after, self.template.approved_file_sha256)
        self.assertNotEqual(generated_content, self.real_template_bytes)
        self.assertTrue(self.template.is_approved)

    def test_long_prescription_keeps_every_medication(self):
        long_medications = [
            Medication.objects.create(
                patient=self.patient,
                prescribed_by=self.dentist,
                name=f"Médicament {index:02d} — essai clinique",
                dosage=f"{index * 10} mg",
                frequency="2 fois par jour matin et soir",
                route="orale",
                start_date=date(2026, 10, 15),
                end_date=date(2026, 11, 15),
                notes=(
                    "Posologie à ajuster selon la tolérance digestive ; "
                    "prendre pendant le repas avec un grand verre d'eau ; "
                    "signaler toute réaction inhabituel au praticien. "
                    f"Référence interne N°{index:03d}."
                ),
            )
            for index in range(1, 41)
        ]
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [medication.pk for medication in long_medications],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201)
        prescription = Prescription.objects.get(pk=response.data["id"])
        with prescription.generated_file.open("rb") as generated:
            document = WordDocument(BytesIO(generated.read()))
        body_text = _document_text(document)

        for medication in long_medications:
            self.assertIn(medication.name, body_text)
        for index in range(1, 41):
            self.assertIn(f"Référence interne N°{index:03d}", body_text)
        self.assertNotIn("{{", body_text)
        # Aucun contenu tronqué : les 40 blocs sont bien délimités.
        self.assertEqual(body_text.count("- Durée : 32 jours"), 40)

    # ------------------------------------------------------------------
    # Cas négatifs : authentification et permissions
    # ------------------------------------------------------------------
    def test_unauthenticated_request_is_rejected(self):
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [self.medications[0].pk],
            },
            format="json",
        )

        self.assertIn(response.status_code, (401, 403))
        self.assertFalse(Prescription.objects.exists())

    def test_non_dentist_role_cannot_generate(self):
        self.client.force_authenticate(self.administrator)
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [self.medications[0].pk],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(Prescription.objects.exists())

    def test_dentist_cannot_generate_for_unlinked_patient(self):
        unlinked_patient = Patient.objects.create(
            first_name="Inconnu",
            last_name="Patient",
            phone="0987654321",
        )
        # Médicament valide pour ce patient et prescrit par le dentiste :
        # la seule faute est l'absence de lien patient/praticien.
        medication = Medication.objects.create(
            patient=unlinked_patient,
            prescribed_by=self.dentist,
            name="Amoxicilline",
            dosage="500 mg",
            frequency="2 fois par jour",
            route="orale",
            start_date=date(2026, 10, 15),
        )
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": unlinked_patient.pk,
                "document_template": self.template.pk,
                "medications": [medication.pk],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(Prescription.objects.exists())

    # ------------------------------------------------------------------
    # Cas négatifs : modèle
    # ------------------------------------------------------------------
    def test_unapproved_real_template_is_rejected(self):
        pending = DocumentTemplate.objects.create(
            name="Ordonnance standard (non approuvée)",
            file=SimpleUploadedFile(
                "ordonnance_standard_pending.docx",
                load_real_template_bytes(),
            ),
        )
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": pending.pk,
                "medications": [self.medications[0].pk],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())

    def test_inactive_real_template_is_rejected(self):
        self.template.is_active = False
        self.template.save(update_fields=("is_active",))
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [self.medications[0].pk],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())

    def test_altered_source_file_is_rejected(self):
        with self.template.file.open("wb") as source:
            source.write(b"source modifiee hors approbation")
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [self.medications[0].pk],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())

    # ------------------------------------------------------------------
    # Cas négatifs : validation des données => aucun document présenté comme valide
    # ------------------------------------------------------------------
    def test_missing_medication_selection_produces_no_artifact(self):
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())
        self.assertFalse(self._generated_files())

    def test_medication_of_another_patient_is_rejected(self):
        other_patient = Patient.objects.create(
            first_name="Autre",
            last_name="Patient",
            phone="0147258369",
        )
        foreign_medication = Medication.objects.create(
            patient=other_patient,
            prescribed_by=self.dentist,
            name="Ibuprofène",
            dosage="400 mg",
            frequency="3 fois par jour",
            route="orale",
            start_date=date(2026, 10, 15),
        )
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [self.medications[0].pk, foreign_medication.pk],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())
        self.assertFalse(self._generated_files())

    def test_medication_prescribed_by_another_dentist_is_rejected(self):
        foreign_medication = Medication.objects.create(
            patient=self.patient,
            prescribed_by=self.other_dentist,
            name="Amoxicilline",
            dosage="250 mg",
            frequency="2 fois par jour",
            route="orale",
            start_date=date(2026, 10, 15),
        )
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [foreign_medication.pk],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())
        self.assertFalse(self._generated_files())

    # ------------------------------------------------------------------
    def _generated_files(self):
        import os

        generated_dir = os.path.join(self.media_dir.name, "prescriptions", "generated")
        if not os.path.isdir(generated_dir):
            return []
        return os.listdir(generated_dir)
