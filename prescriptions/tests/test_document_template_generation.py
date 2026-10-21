from tempfile import TemporaryDirectory
from datetime import timedelta
from io import BytesIO

from docx import Document as WordDocument
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.utils import timezone
from rest_framework.test import APITestCase

from accounts.models import User
from appointments.models import Appointment
from documents.models import DocumentTemplate
from documents.tests.docx_fixtures import make_prescription_docx, make_simple_docx
from patients.models import Medication, Patient
from prescriptions.models import Prescription, PrescriptionTemplate


class PrescriptionDocumentTemplateGenerationTests(APITestCase):
    def setUp(self):
        self.media_dir = TemporaryDirectory()
        self.addCleanup(self.media_dir.cleanup)
        self.settings_context = override_settings(MEDIA_ROOT=self.media_dir.name)
        self.settings_context.enable()
        self.addCleanup(self.settings_context.disable)

        self.dentist = User.objects.create_user(
            email="dentist@example.test",
            password="test-password",
            first_name="Dental",
            last_name="Doctor",
            role=User.Role.DENTIST,
        )
        self.administrator = User.objects.create_user(
            email="admin@example.test",
            password="test-password",
            first_name="Admin",
            last_name="User",
            role=User.Role.ADMINISTRATOR,
        )
        self.patient = Patient.objects.create(
            first_name="Patient",
            last_name="Example",
            phone="0123456789",
        )
        starts = timezone.now() + timedelta(days=1)
        Appointment.objects.create(
            patient=self.patient,
            practitioner=self.dentist,
            start_at=starts,
            end_at=starts + timedelta(minutes=30),
        )
        self.medications = [
            Medication.objects.create(
                patient=self.patient,
                prescribed_by=self.dentist,
                name="Amoxicilline",
                dosage="500 mg",
                frequency="3 fois par jour",
                route="orale",
                start_date=timezone.localdate(),
                notes="Après le repas",
            ),
            Medication.objects.create(
                patient=self.patient,
                prescribed_by=self.dentist,
                name="Paracétamol",
                dosage="1 g",
                frequency="si douleur",
                route="orale",
                start_date=timezone.localdate(),
                notes="بعد الوجبة",
            ),
        ]
        self.template = DocumentTemplate.objects.create(
            name="Prescription",
            file=SimpleUploadedFile(
                "prescription.docx",
                make_prescription_docx(),
            ),
        )
        self.client.force_authenticate(self.administrator)
        approval_response = self.client.post(
            f"/api/document-templates/{self.template.pk}/approve/",
            format="json",
        )
        self.assertEqual(approval_response.status_code, 200)

    def test_generation_copies_source_and_edit_does_not_change_it(self):
        self.client.force_authenticate(self.dentist)

        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [medication.pk for medication in self.medications],
                "notes": "Prendre après le repas",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201)
        prescription = Prescription.objects.get(pk=response.data["id"])
        self.assertEqual(prescription.document_template_id, self.template.pk)
        self.assertIsNone(prescription.template_id)
        self.assertIn(
            "Amoxicilline",
            prescription.filled_data["medications"][0]["name"],
        )
        with prescription.generated_file.open("rb") as generated:
            generated_content = generated.read()
        parsed = WordDocument(BytesIO(generated_content))
        body_text = "\n".join(
            [paragraph.text for paragraph in parsed.paragraphs]
            + [
                cell.text
                for table in parsed.tables
                for row in table.rows
                for cell in row.cells
            ]
        )
        self.assertIn("Patient : Patient Example", body_text)
        self.assertIn("Paracétamol", body_text)
        self.assertIn("بعد الوجبة", body_text)
        self.assertIn("Signature", body_text)
        self.assertNotIn("{{", body_text)
        self.assertEqual(len(parsed.tables[0].rows), 2)
        header_text = parsed.sections[0].header.paragraphs[0].text
        self.assertEqual(
            header_text,
            "Praticien : Dental Doctor",
        )
        footer_text = parsed.sections[0].footer.paragraphs[0].text
        self.assertIn("Prendre après le repas", footer_text)
        self.assertIn(
            prescription.generated_at.strftime("%d/%m/%Y"),
            footer_text,
        )
        self.assertNotIn("{{", footer_text)
        with self.template.file.open("rb") as source:
            source_content = source.read()
        self.assertNotEqual(generated_content, source_content)

        replacement_source_content = make_prescription_docx()
        self.template.file = SimpleUploadedFile(
            "prescription.docx",
            replacement_source_content,
        )
        self.template.save()
        self.template.refresh_from_db()
        self.assertFalse(self.template.is_approved)
        with prescription.generated_file.open("rb") as generated:
            self.assertEqual(generated.read(), generated_content)

        edit_response = self.client.patch(
            f"/api/prescriptions/{prescription.pk}/",
            {
                "generated_file": SimpleUploadedFile(
                    "edited-prescription.docx",
                    make_simple_docx(),
                )
            },
            format="multipart",
        )

        self.assertEqual(edit_response.status_code, 200)
        prescription.refresh_from_db()
        with self.template.file.open("rb") as source:
            self.assertEqual(source.read(), replacement_source_content)
        with prescription.generated_file.open("rb") as generated:
            edited = WordDocument(BytesIO(generated.read()))
        self.assertIn("Modèle administratif", edited.paragraphs[0].text)

    def test_generation_rejects_unapproved_template(self):
        unapproved = DocumentTemplate.objects.create(
            name="Pending",
            file=SimpleUploadedFile("pending.docx", make_simple_docx()),
        )
        self.client.force_authenticate(self.dentist)

        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {"patient": self.patient.pk, "document_template": unapproved.pk},
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())

    def test_dentist_cannot_bypass_approval_with_legacy_template_endpoint(self):
        legacy_template = PrescriptionTemplate.objects.create(
            name="Legacy",
            description="",
            content="legacy template",
        )
        self.client.force_authenticate(self.dentist)

        response = self.client.post(
            "/api/prescriptions/",
            {
                "patient": self.patient.pk,
                "dentist": self.dentist.pk,
                "template": legacy_template.pk,
                "status": Prescription.Status.DRAFT,
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())

    def test_generation_rejects_missing_medicine_selection(self):
        self.client.force_authenticate(self.dentist)

        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {"patient": self.patient.pk, "document_template": self.template.pk},
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())

    def test_generation_rejects_inactive_template(self):
        self.template.is_active = False
        self.template.save(update_fields=("is_active",))
        self.client.force_authenticate(self.dentist)

        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [self.medications[0].pk],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())

    def test_generation_rejects_source_changed_without_model_save(self):
        with self.template.file.open("rb") as source:
            original = source.read()
        changed_source = make_simple_docx("Source modifiée")
        with self.template.file.open("wb") as source:
            source.write(changed_source)
        self.assertNotEqual(original, changed_source)

        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [self.medications[0].pk],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())

    def test_generation_rejects_missing_source_file(self):
        self.template.file.storage.delete(self.template.file.name)
        self.client.force_authenticate(self.dentist)

        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [self.medications[0].pk],
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())

    def test_generation_rejects_patient_not_linked_to_dentist(self):
        other_dentist = User.objects.create_user(
            email="other@example.test",
            password="test-password",
            first_name="Other",
            last_name="Doctor",
            role=User.Role.DENTIST,
        )
        other_patient = Patient.objects.create(
            first_name="Other",
            last_name="Patient",
            phone="0987654321",
        )
        starts = timezone.now() + timedelta(days=2)
        Appointment.objects.create(
            patient=other_patient,
            practitioner=other_dentist,
            start_at=starts,
            end_at=starts + timedelta(minutes=30),
        )
        self.client.force_authenticate(self.dentist)

        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {"patient": other_patient.pk, "document_template": self.template.pk},
            format="json",
        )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(Prescription.objects.exists())
