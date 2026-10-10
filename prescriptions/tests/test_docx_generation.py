from datetime import date
from io import BytesIO
import json
from pathlib import Path
import tempfile
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import User
from documents.docx import DocxTemplateError, render_docx
from documents.storage import private_document_storage
from patients.models import Medication, Patient
from prescriptions.models import Prescription, PrescriptionTemplate
from prescriptions.services import build_prescription_context


class PrescriptionDocxGenerationTests(APITestCase):
    def setUp(self):
        self.temp_storage = tempfile.TemporaryDirectory(prefix='prescription-docx-')
        self.storage_override = override_settings(
            PRIVATE_DOCUMENT_ROOT=self.temp_storage.name,
        )
        self.storage_override.enable()

        self.dentist = User.objects.create_user(
            email='dentist@example.test',
            password='test-password',
            first_name='Amine',
            last_name='Benali',
            role=User.Role.DENTIST,
        )
        self.other_dentist = User.objects.create_user(
            email='other@example.test',
            password='test-password',
            first_name='Autre',
            last_name='Dentiste',
            role=User.Role.DENTIST,
        )
        self.patient = Patient.objects.create(
            first_name='Sofia',
            last_name='Martin',
            phone='0600000000',
            birthdate=date(1990, 5, 10),
        )

        template_bytes = self.make_template()
        self.template = PrescriptionTemplate.objects.create(
            name='Ordonnance DOCX',
            description='Template de test',
            content='',
            variables=[
                {
                    'name': 'patient',
                    'data_type': 'object',
                    'is_required': True,
                    'properties': {
                        'first_name': {'data_type': 'string', 'is_required': True},
                        'last_name': {'data_type': 'string', 'is_required': True},
                    },
                },
                {
                    'name': 'dentist',
                    'data_type': 'object',
                    'is_required': True,
                    'properties': {
                        'first_name': {'data_type': 'string', 'is_required': True},
                        'last_name': {'data_type': 'string', 'is_required': True},
                    },
                },
                {
                    'name': 'medications',
                    'data_type': 'list',
                    'is_required': True,
                    'allow_empty': False,
                    'items': {
                        'data_type': 'object',
                        'properties': {
                            'name': {'data_type': 'string', 'is_required': True},
                            'dosage': {'data_type': 'string', 'is_required': True},
                        },
                    },
                },
                {'name': 'generated_at', 'data_type': 'string', 'is_required': True},
            ],
            docx_file=SimpleUploadedFile(
                'prescription.docx',
                template_bytes,
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            ),
            is_active=True,
        )
        self.prescription = Prescription.objects.create(
            patient=self.patient,
            dentist=self.dentist,
            template=self.template,
            filled_data={
                'medications': [
                    {'name': 'Amoxicilline', 'dosage': '500 mg'},
                    {'name': 'Ibuprofène', 'dosage': '200 mg'},
                ],
            },
            status=Prescription.Status.DRAFT,
        )
        self.client.force_authenticate(self.dentist)

    def tearDown(self):
        self.storage_override.disable()
        self.temp_storage.cleanup()
        super().tearDown()

    @staticmethod
    def make_template(unknown_placeholder=False):
        document = Document()
        document.add_paragraph('Patient : {{ patient.first_name }} {{ patient.last_name }}')
        if unknown_placeholder:
            document.add_paragraph('{{ patient.unknown_field }}')
        document.add_paragraph('Dentiste : {{ dentist.first_name }} {{ dentist.last_name }}')
        document.add_paragraph('Date : {{ generated_at }}')
        arabic_paragraph = document.add_paragraph('معلومات المريض')
        arabic_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        document.add_paragraph('{%p for medication in medications %}')
        document.add_paragraph('{{ medication.name }} | {{ medication.dosage }}')
        document.add_paragraph('{%p endfor %}')
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    def test_generation_and_private_download_return_a_real_docx(self):
        response = self.client.post(
            reverse('prescription-generate-docx', kwargs={'pk': self.prescription.pk}),
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.prescription.refresh_from_db()
        self.assertEqual(self.prescription.status, Prescription.Status.GENERATED)
        self.assertTrue(self.prescription.generated_docx.name)
        self.assertFalse(response.data['docx_download_url'].startswith('/media/'))
        self.assertTrue(self.prescription.generated_docx.path.startswith(self.temp_storage.name))
        self.assertNotIn('docx_file', response.data)
        self.assertNotIn('generated_docx', response.data)
        with self.assertRaises(ValueError):
            self.prescription.generated_docx.url

        download = self.client.get(
            reverse('prescription-download-docx', kwargs={'pk': self.prescription.pk}),
        )
        self.assertEqual(download.status_code, status.HTTP_200_OK)
        self.assertEqual(
            download['Content-Type'],
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        self.assertIn('.docx', download['Content-Disposition'])

        downloaded = b''.join(download.streaming_content)
        rendered = Document(BytesIO(downloaded))
        text = '\n'.join(paragraph.text for paragraph in rendered.paragraphs)
        self.assertIn('Patient : Sofia Martin', text)
        self.assertIn('Dentiste : Amine Benali', text)
        self.assertIn('معلومات المريض', text)
        self.assertIn(
            f"Date : {timezone.localtime(self.prescription.generated_at).strftime('%d/%m/%Y')}",
            text,
        )
        self.assertIn('Amoxicilline | 500 mg', text)
        self.assertIn('Ibuprofène | 200 mg', text)
        self.assertEqual(text.count('Amoxicilline'), 1)
        self.assertEqual(text.count('Ibuprofène'), 1)

    def test_required_variable_missing_returns_explicit_error(self):
        self.prescription.filled_data.pop('medications')
        self.prescription.save(update_fields=['filled_data'])

        response = self.client.post(
            reverse('prescription-generate-docx', kwargs={'pk': self.prescription.pk}),
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('medications', response.data['detail'])
        self.assertFalse(self.prescription.generated_docx)

    def test_empty_medication_list_is_rejected_by_template_contract(self):
        self.prescription.filled_data['medications'] = []
        self.prescription.save(update_fields=['filled_data'])

        response = self.client.post(
            reverse('prescription-generate-docx', kwargs={'pk': self.prescription.pk}),
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('ne peut pas être vide', response.data['detail'])

    def test_unknown_docx_placeholder_is_rejected_when_template_is_saved(self):
        invalid_docx = self.make_template(unknown_placeholder=True)
        serializer_response = self.client.post(
            reverse('prescriptiontemplate-list'),
            {
                'name': 'Template invalide',
                'description': '',
                'content': 'Ordonnance de test',
                'variables': json.dumps(self.template.variables),
                'is_active': True,
                'docx_file': SimpleUploadedFile(
                    'unknown.docx',
                    invalid_docx,
                    content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                ),
            },
            format='multipart',
        )

        self.assertEqual(serializer_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('patient.unknown_field', str(serializer_response.data))

    def test_invalid_docx_is_rejected_when_template_is_saved(self):
        response = self.client.post(
            reverse('prescriptiontemplate-list'),
            {
                'name': 'Fichier invalide',
                'description': '',
                'content': 'Ordonnance de test',
                'variables': '[]',
                'is_active': True,
                'docx_file': SimpleUploadedFile(
                    'invalid.docx',
                    b'not-a-docx',
                    content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                ),
            },
            format='multipart',
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('illisible ou invalide', str(response.data))

    def test_template_upload_is_validated_and_never_returns_a_file_path(self):
        response = self.client.post(
            reverse('prescriptiontemplate-list'),
            {
                'name': 'Template uploadé',
                'description': '',
                'content': 'Ordonnance de test',
                'variables': json.dumps(self.template.variables),
                'is_active': True,
                'docx_file': SimpleUploadedFile(
                    'valid.docx',
                    self.make_template(),
                    content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                ),
            },
            format='multipart',
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertNotIn('docx_file', response.data)
        template = PrescriptionTemplate.objects.get(name='Template uploadé')
        self.assertTrue(template.docx_file.path.startswith(self.temp_storage.name))
        with self.assertRaises(ValueError):
            template.docx_file.url

    def test_malformed_repeating_block_is_rejected_when_template_is_saved(self):
        invalid_template = Document()
        invalid_template.add_paragraph('{%p for medication in medications %}')
        invalid_template.add_paragraph('{{ medication.name }}')
        output = BytesIO()
        invalid_template.save(output)

        response = self.client.post(
            reverse('prescriptiontemplate-list'),
            {
                'name': 'Bloc répétable invalide',
                'description': '',
                'content': 'Ordonnance de test',
                'variables': json.dumps(self.template.variables),
                'is_active': True,
                'docx_file': SimpleUploadedFile(
                    'malformed.docx',
                    output.getvalue(),
                    content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                ),
            },
            format='multipart',
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('mal formés', str(response.data))

    def test_incompatible_medication_item_type_returns_explicit_error(self):
        self.prescription.filled_data['medications'] = ['not-an-object']
        self.prescription.save(update_fields=['filled_data'])

        response = self.client.post(
            reverse('prescription-generate-docx', kwargs={'pk': self.prescription.pk}),
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('incompatible avec le type object', response.data['detail'])

    def test_optional_variable_uses_its_declared_default(self):
        template = Document()
        template.add_paragraph('Note : {{ note }}')
        output = BytesIO()
        template.save(output)
        rendered_bytes = render_docx(
            output.getvalue(),
            [{
                'name': 'note',
                'data_type': 'string',
                'is_required': False,
                'default_value': 'Aucune note',
            }],
            {},
        )

        rendered = Document(BytesIO(rendered_bytes))
        self.assertIn('Note : Aucune note', '\n'.join(p.text for p in rendered.paragraphs))

    def test_empty_list_is_allowed_when_template_contract_says_so(self):
        template = Document()
        template.add_paragraph('{%p for medication in medications %}')
        template.add_paragraph('{{ medication.name }}')
        template.add_paragraph('{%p endfor %}')
        output = BytesIO()
        template.save(output)
        rendered_bytes = render_docx(
            output.getvalue(),
            [{
                'name': 'medications',
                'data_type': 'list',
                'is_required': True,
                'allow_empty': True,
                'items': {
                    'data_type': 'object',
                    'properties': {
                        'name': {'data_type': 'string', 'is_required': True},
                    },
                },
            }],
            {'medications': []},
        )

        rendered = Document(BytesIO(rendered_bytes))
        self.assertEqual(rendered.paragraphs, [])

    def test_non_owner_dentist_cannot_generate_or_download_another_dentists_prescription(self):
        generated = self.client.post(
            reverse('prescription-generate-docx', kwargs={'pk': self.prescription.pk}),
        )
        self.assertEqual(generated.status_code, status.HTTP_200_OK)

        self.client.force_authenticate(self.other_dentist)

        generate = self.client.post(
            reverse('prescription-generate-docx', kwargs={'pk': self.prescription.pk}),
        )
        download = self.client.get(
            reverse('prescription-download-docx', kwargs={'pk': self.prescription.pk}),
        )

        self.assertEqual(generate.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(download.status_code, status.HTTP_403_FORBIDDEN)

        prescription_list = self.client.get(reverse('prescription-list'))
        serialized = next(item for item in prescription_list.data if item['id'] == self.prescription.pk)
        self.assertIsNone(serialized['docx_download_url'])

    def test_admin_template_page_does_not_require_a_public_file_url(self):
        administrator = User.objects.create_superuser(
            email='admin@example.test',
            password='admin-password',
            first_name='Admin',
            last_name='Test',
        )
        self.client.force_authenticate(user=None)
        self.client.force_login(administrator)

        response = self.client.get(
            reverse('admin:prescriptions_prescriptiontemplate_change', args=[self.template.pk]),
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_prescription_context_uses_existing_business_sources_without_mutating_them(self):
        medication = Medication.objects.create(
            patient=self.patient,
            name='Traitement habituel',
            dosage='10 mg',
            frequency='1 fois/jour',
            route='orale',
            start_date=date(2026, 1, 1),
            prescribed_by=self.dentist,
        )
        context, generated_at = build_prescription_context(self.prescription)

        self.assertEqual(context['patient']['patient_code'], self.patient.patient_code)
        self.assertEqual(context['dentist']['email'], self.dentist.email)
        self.assertEqual(context['medications'], self.prescription.filled_data['medications'])
        self.assertEqual(Medication.objects.get(pk=medication.pk).patient_id, self.patient.pk)
        self.assertIsNotNone(generated_at.tzinfo)

    def test_invalid_zip_docx_is_rejected_by_prevalidation(self):
        with self.assertRaises(DocxTemplateError):
            render_docx(b'PK\x03\x04invalid', self.template.variables, {})
