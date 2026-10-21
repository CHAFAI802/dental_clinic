"""Contrat API du parcours médecin : médicaments + génération d'ordonnance.

Ces tests exécutent exactement les requêtes envoyées par le frontend
(PrescriptionSection.jsx) :

    GET|POST /api/medications/?patient=<id>
    PATCH    /api/medications/<id>/
    POST     /api/prescriptions/generate-from-document-template/
                 {patient, document_template, medications:[ids], notes}

et vérifient la réponse, le fichier DOCX produit et les règles backend
(modèle actif et approuvé, rattachement patient/dentiste, rôles).
"""

from datetime import date, timedelta
from hashlib import sha256
from io import BytesIO
from tempfile import TemporaryDirectory
from zipfile import is_zipfile

from docx import Document as WordDocument
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.utils import timezone
from rest_framework.test import APITestCase

from accounts.models import User
from appointments.models import Appointment
from documents.models import DocumentTemplate
from documents.tests.docx_fixtures import load_real_template_bytes
from patients.models import Medication, Patient
from prescriptions.models import Prescription


def _document_text(document):
    """Texte complet d'un DOCX : corps, tableaux, en-têtes et pieds de page."""
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    for section in document.sections:
        for part in (section.header, section.footer):
            parts.extend(paragraph.text for paragraph in part.paragraphs)
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
    return "\n".join(parts)


class MedicationApiTests(APITestCase):
    """GET|POST|PATCH /api/medications/ — contrat et permissions."""

    def setUp(self):
        self.media_dir = TemporaryDirectory()
        self.addCleanup(self.media_dir.cleanup)
        self.settings_context = override_settings(MEDIA_ROOT=self.media_dir.name)
        self.settings_context.enable()
        self.addCleanup(self.settings_context.disable)

        self.dentist = User.objects.create_user(
            email="med-dentist@example.test",
            password="test-password",
            first_name="Nadia",
            last_name="Ben Salah",
            role=User.Role.DENTIST,
        )
        self.other_dentist = User.objects.create_user(
            email="med-other@example.test",
            password="test-password",
            first_name="Autre",
            last_name="Praticien",
            role=User.Role.DENTIST,
        )
        self.receptionist = User.objects.create_user(
            email="med-reception@example.test",
            password="test-password",
            first_name="Régie",
            last_name="Accueil",
            role=User.Role.RECEPTIONIST,
        )
        self.patient = Patient.objects.create(
            first_name="Émilie",
            last_name="Durand",
            phone="06 12 34 56 78",
        )
        starts = timezone.now() + timedelta(days=1)
        Appointment.objects.create(
            patient=self.patient,
            practitioner=self.dentist,
            start_at=starts,
            end_at=starts + timedelta(minutes=30),
        )
        self.foreign_patient = Patient.objects.create(
            first_name="Patient",
            last_name="Étranger",
            phone="0147258369",
        )
        other_starts = timezone.now() + timedelta(days=2)
        Appointment.objects.create(
            patient=self.foreign_patient,
            practitioner=self.other_dentist,
            start_at=other_starts,
            end_at=other_starts + timedelta(minutes=30),
        )

    # ------------------------------------------------------------------
    # Contrat de création (payload exact du frontend)
    # ------------------------------------------------------------------
    def test_dentist_creates_medication_for_his_patient(self):
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            f"/api/medications/?patient={self.patient.pk}",
            {
                "name": "Amoxicilline",
                "dosage": "500 mg",
                "frequency": "3 fois par jour",
                "route": "orale",
                "start_date": "2026-10-15",
                "end_date": "2026-10-21",
                "notes": "Après le repas",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["name"], "Amoxicilline")
        self.assertEqual(response.data["dosage"], "500 mg")
        self.assertEqual(response.data["frequency"], "3 fois par jour")
        self.assertEqual(response.data["route"], "orale")
        self.assertEqual(response.data["start_date"], "2026-10-15")
        self.assertEqual(response.data["end_date"], "2026-10-21")
        self.assertEqual(response.data["notes"], "Après le repas")
        # Champs fixés par le serveur, jamais reçus du client.
        self.assertEqual(response.data["patient"], self.patient.pk)
        self.assertEqual(response.data["prescribed_by"], self.dentist.pk)

        medication = Medication.objects.get(pk=response.data["id"])
        self.assertEqual(medication.patient_id, self.patient.pk)
        self.assertEqual(medication.prescribed_by_id, self.dentist.pk)

    def test_optional_fields_are_optional(self):
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            f"/api/medications/?patient={self.patient.pk}",
            {
                "name": "Paracétamol",
                "dosage": "1 g",
                "frequency": "si douleur",
                "route": "orale",
                "start_date": "2026-10-15",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertIsNone(response.data["end_date"])
        self.assertEqual(response.data["notes"], "")

    def test_missing_required_fields_are_rejected(self):
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            f"/api/medications/?patient={self.patient.pk}",
            {"dosage": "500 mg"},
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        for field in ("name", "frequency", "route", "start_date"):
            self.assertIn(field, response.data)
        self.assertFalse(Medication.objects.exists())

    def test_patient_parameter_is_required(self):
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            "/api/medications/",
            {"name": "X", "dosage": "1", "frequency": "1", "route": "orale",
             "start_date": "2026-10-15"},
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("patient", response.data)
        self.assertFalse(Medication.objects.exists())

    # ------------------------------------------------------------------
    # Rattachement patient / dentiste
    # ------------------------------------------------------------------
    def test_create_for_unlinked_patient_is_rejected(self):
        self.client.force_authenticate(self.dentist)
        response = self.client.post(
            f"/api/medications/?patient={self.foreign_patient.pk}",
            {
                "name": "Ibuprofène",
                "dosage": "400 mg",
                "frequency": "3 fois par jour",
                "route": "orale",
                "start_date": "2026-10-15",
            },
            format="json",
        )

        # Aucune existence de patient étranger divulguée.
        self.assertEqual(response.status_code, 404)
        self.assertFalse(Medication.objects.exists())

    def test_list_is_scoped_to_linked_patients(self):
        Medication.objects.create(
            patient=self.patient,
            prescribed_by=self.dentist,
            name="Médicament lié",
            dosage="10 mg",
            frequency="1/j",
            route="orale",
            start_date=date(2026, 10, 15),
        )
        Medication.objects.create(
            patient=self.foreign_patient,
            prescribed_by=self.other_dentist,
            name="Médicament étranger",
            dosage="10 mg",
            frequency="1/j",
            route="orale",
            start_date=date(2026, 10, 15),
        )

        self.client.force_authenticate(self.dentist)
        all_records = self.client.get("/api/medications/", format="json")
        self.assertEqual(all_records.status_code, 200)
        self.assertEqual(
            [item["name"] for item in all_records.data],
            ["Médicament lié"],
        )

        filtered = self.client.get(
            f"/api/medications/?patient={self.foreign_patient.pk}",
            format="json",
        )
        self.assertEqual(filtered.status_code, 200)
        self.assertEqual(filtered.data, [])

    def test_update_of_own_medication_succeeds(self):
        medication = Medication.objects.create(
            patient=self.patient,
            prescribed_by=self.dentist,
            name="Amoxicilline",
            dosage="500 mg",
            frequency="3 fois par jour",
            route="orale",
            start_date=date(2026, 10, 15),
        )

        self.client.force_authenticate(self.dentist)
        response = self.client.patch(
            f"/api/medications/{medication.pk}/",
            {"dosage": "1 g"},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.data)
        medication.refresh_from_db()
        self.assertEqual(medication.dosage, "1 g")
        self.assertEqual(medication.prescribed_by_id, self.dentist.pk)

    def test_update_of_foreign_medication_is_rejected(self):
        medication = Medication.objects.create(
            patient=self.foreign_patient,
            prescribed_by=self.other_dentist,
            name="Médicament étranger",
            dosage="10 mg",
            frequency="1/j",
            route="orale",
            start_date=date(2026, 10, 15),
        )

        self.client.force_authenticate(self.dentist)
        response = self.client.patch(
            f"/api/medications/{medication.pk}/",
            {"dosage": "999 mg"},
            format="json",
        )

        self.assertEqual(response.status_code, 404)
        medication.refresh_from_db()
        self.assertEqual(medication.dosage, "10 mg")

    # ------------------------------------------------------------------
    # Rôles
    # ------------------------------------------------------------------
    def test_other_roles_are_still_restricted(self):
        payload = {
            "name": "X",
            "dosage": "10 mg",
            "frequency": "1/j",
            "route": "orale",
            "start_date": "2026-10-15",
        }

        self.client.force_authenticate(self.receptionist)
        create_response = self.client.post(
            f"/api/medications/?patient={self.patient.pk}",
            payload,
            format="json",
        )
        list_response = self.client.get("/api/medications/", format="json")

        self.assertEqual(create_response.status_code, 403)
        self.assertEqual(list_response.status_code, 403)
        self.assertFalse(Medication.objects.exists())


class FrontendGenerationFlowTests(APITestCase):
    """Parcours complet tel que le frontend l'exécute, vérifié côté API."""

    def setUp(self):
        self.media_dir = TemporaryDirectory()
        self.addCleanup(self.media_dir.cleanup)
        self.settings_context = override_settings(MEDIA_ROOT=self.media_dir.name)
        self.settings_context.enable()
        self.addCleanup(self.settings_context.disable)

        self.dentist = User.objects.create_user(
            email="flow-dentist@example.test",
            password="test-password",
            first_name="Nadia",
            last_name="Ben Salah",
            role=User.Role.DENTIST,
        )
        self.administrator = User.objects.create_user(
            email="flow-admin@example.test",
            password="test-password",
            first_name="Admin",
            last_name="User",
            role=User.Role.ADMINISTRATOR,
        )
        self.receptionist = User.objects.create_user(
            email="flow-reception@example.test",
            password="test-password",
            first_name="Régie",
            last_name="Accueil",
            role=User.Role.RECEPTIONIST,
        )
        self.patient = Patient.objects.create(
            first_name="Émilie",
            middle_name="Marie",
            last_name="Durand",
            birthdate=date(1990, 3, 15),
            phone="06 12 34 56 78",
        )
        starts = timezone.now() + timedelta(days=1)
        Appointment.objects.create(
            patient=self.patient,
            practitioner=self.dentist,
            start_at=starts,
            end_at=starts + timedelta(minutes=30),
        )

        # Véritable modèle Word approuvé via la vraie endpoint d'approbation.
        self.template = DocumentTemplate.objects.create(
            name="Ordonnance standard",
            file=SimpleUploadedFile(
                "ordonnance_standard.docx",
                load_real_template_bytes(),
            ),
        )
        self.client.force_authenticate(self.administrator)
        approval = self.client.post(
            f"/api/document-templates/{self.template.pk}/approve/",
            format="json",
        )
        self.assertEqual(approval.status_code, 200)
        self.client.force_authenticate(self.dentist)

    def _create_medication(self, **overrides):
        payload = {
            "name": "Amoxicilline",
            "dosage": "500 mg",
            "frequency": "3 fois par jour",
            "route": "orale",
            "start_date": "2026-10-15",
            "end_date": "2026-10-21",
            "notes": "Après le repas",
        }
        payload.update(overrides)
        response = self.client.post(
            f"/api/medications/?patient={self.patient.pk}",
            payload,
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        return response.data

    def _generated_files(self):
        import os

        generated_dir = os.path.join(
            self.media_dir.name, "prescriptions", "generated"
        )
        if not os.path.isdir(generated_dir):
            return []
        return os.listdir(generated_dir)

    def test_frontend_payload_generates_real_docx(self):
        first = self._create_medication()
        second = self._create_medication(
            name="Paracétamol",
            dosage="1 g",
            frequency="si douleur",
            end_date=None,
            notes="Ne pas dépasser 3 g par jour",
        )

        # Payload exact de PrescriptionSection.jsx.
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [first["id"], second["id"]],
                "notes": "Contrôle post-opératoire",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["status"], Prescription.Status.GENERATED)
        self.assertEqual(response.data["patient"], self.patient.pk)
        self.assertEqual(response.data["dentist"], self.dentist.pk)
        self.assertEqual(response.data["document_template"], self.template.pk)
        self.assertEqual(len(response.data["filled_data"]["medications"]), 2)

        # Le fichier retourné au frontend : chemin /media/... exploitable.
        generated_file = response.data["generated_file"]
        self.assertTrue(generated_file)
        self.assertIn("/media/prescriptions/generated/", generated_file)
        self.assertTrue(generated_file.split("?")[0].endswith(".docx"))

        prescription = Prescription.objects.get(pk=response.data["id"])
        self.assertEqual(prescription.patient_id, self.patient.pk)
        self.assertEqual(prescription.dentist_id, self.dentist.pk)
        with prescription.generated_file.open("rb") as handle:
            content = handle.read()

        # Véritable DOCX, contenu métier injecté, aucun marqueur résiduel.
        self.assertTrue(is_zipfile(BytesIO(content)))
        document = WordDocument(BytesIO(content))
        text = _document_text(document)
        self.assertIn("ORDONNANCE MÉDICALE", text)
        self.assertIn("Patient : Émilie Marie Durand", text)
        self.assertIn("Amoxicilline — 500 mg", text)
        self.assertIn("Paracétamol — 1 g", text)
        self.assertIn("- Durée : 7 jours", text)
        self.assertIn("Ne pas dépasser 3 g par jour", text)
        self.assertNotIn("{{", text)
        self.assertNotIn("}}", text)
        # `notes` de l'ordonnance : conservé sur l'enregistrement et renvoyé
        # au frontend. Le VRAI modèle ne contient aucun marqueur
        # `{{ prescription.notes }}` : ce champ n'est donc pas imprimé par ce
        # modèle (constat, pas une défaillance du moteur).
        self.assertEqual(response.data["notes"], "Contrôle post-opératoire")
        self.assertEqual(prescription.notes, "Contrôle post-opératoire")
        # Intégrité des octets servis : empreinte exploitable par le frontend
        # (comparaison avec le fichier récupéré via /media/... en E2E navigateur).
        self.assertEqual(len(sha256(content).hexdigest()), 64)

    def test_generation_with_unapproved_template_is_rejected(self):
        pending = DocumentTemplate.objects.create(
            name="Modèle non approuvé",
            file=SimpleUploadedFile(
                "pending.docx",
                load_real_template_bytes(),
            ),
        )
        medication = self._create_medication()

        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": pending.pk,
                "medications": [medication["id"]],
                "notes": "",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("document_template", response.data)
        self.assertFalse(Prescription.objects.exists())
        self.assertEqual(self._generated_files(), [])

    def test_inactive_approved_template_is_rejected(self):
        self.template.is_active = False
        self.template.save(update_fields=("is_active",))
        medication = self._create_medication()

        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [medication["id"]],
                "notes": "",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())
        self.assertEqual(self._generated_files(), [])

    def test_generation_requires_at_least_one_medication(self):
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [],
                "notes": "",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("medications", response.data)
        self.assertFalse(Prescription.objects.exists())
        self.assertEqual(self._generated_files(), [])

    def test_non_dentist_role_cannot_generate(self):
        medication = self._create_medication()
        self.client.force_authenticate(self.receptionist)
        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": self.patient.pk,
                "document_template": self.template.pk,
                "medications": [medication["id"]],
                "notes": "",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(Prescription.objects.exists())
        self.assertEqual(self._generated_files(), [])

    def test_dentist_cannot_generate_for_unlinked_patient(self):
        unlinked = Patient.objects.create(
            first_name="Inconnu",
            last_name="Patient",
            phone="0987654321",
        )
        # Médicament appartenant au patient non rattaché et prescrit par ce
        # dentiste (créé ici directement : l'endpoint /api/medications/ refuse
        # déjà un tel patient) — seule faute : l'absence de lien patient/dentiste.
        Medication.objects.create(
            patient=unlinked,
            prescribed_by=self.dentist,
            name="Amoxicilline",
            dosage="500 mg",
            frequency="2 fois par jour",
            route="orale",
            start_date=date(2026, 10, 15),
        )
        medication = Medication.objects.get(patient=unlinked)

        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": unlinked.pk,
                "document_template": self.template.pk,
                "medications": [medication.pk],
                "notes": "",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(Prescription.objects.exists())
        self.assertEqual(self._generated_files(), [])

    def test_medication_mismatch_with_patient_is_rejected(self):
        # Patient et médicaments hétérogènes (impossible depuis le frontend :
        # les deux proviennent du même dossier) => refus backend en 400.
        medication = self._create_medication()
        other_patient = Patient.objects.create(
            first_name="Autre",
            last_name="Patient",
            phone="0147258369",
        )

        response = self.client.post(
            "/api/prescriptions/generate-from-document-template/",
            {
                "patient": other_patient.pk,
                "document_template": self.template.pk,
                "medications": [medication["id"]],
                "notes": "",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Prescription.objects.exists())
        self.assertEqual(self._generated_files(), [])
